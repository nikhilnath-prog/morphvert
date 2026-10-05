"""Pure PDF and image conversion helpers for the Morphvert MVP."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException, status


def ensure_pdf_bytes(filename: str, content_type: str, content: bytes) -> None:
    """Validate an uploaded file as a PDF payload."""
    if Path(filename).suffix.lower() != ".pdf":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only PDF files are supported")
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded PDF is empty")

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="pypdf is not installed") from exc

    try:
        PdfReader(BytesIO(content))
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Corrupted PDF file") from exc


def ensure_excel_bytes(filename: str, content_type: str, content: bytes) -> None:
    """Validate an uploaded file as an XLSX payload."""
    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only XLSX files are supported")
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded Excel file is empty")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="openpyxl is not installed") from exc

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if not workbook.sheetnames:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel workbook is empty")
        workbook.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Corrupted Excel file") from exc


def convert_pdf_to_docx(input_pdf: Path, output_docx: Path) -> None:
    print("STEP 1: Starting PDF to DOCX conversion")

    try:
        from pdf2docx import Converter
        print("STEP 2: pdf2docx imported successfully")
    except Exception as exc:
        print(f"STEP 2 FAILED: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=f"pdf2docx import failed: {str(exc)}",
        )

    try:
        converter = Converter(str(input_pdf))
        print("STEP 3: Converter created")

        converter.convert(str(output_docx), start=0, end=None)
        print("STEP 4: Conversion completed")

    except Exception as exc:
        print(f"STEP 4 FAILED: {exc}")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Conversion failed: {str(exc)}",
        )

    finally:
        try:
            converter.close()
        except:
            pass


def convert_pdf_to_excel(input_pdf: Path, output_xlsx: Path) -> None:
    """Convert a PDF into a simple XLSX workbook by extracting text rows."""
    try:
        from openpyxl import Workbook
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="openpyxl or pypdf is not installed") from exc

    reader = PdfReader(str(input_pdf))
    if not reader.pages:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded PDF has no pages")

    workbook = Workbook()
    default_sheet = workbook.active
    default_sheet.title = "Sheet1"
    extracted_rows: list[list[str]] = []
    has_table_structure = False

    for page_index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or page.extract_text(extraction_mode="layout") or ""
        if not page_text.strip():
            continue

        lines = [line.strip() for line in page_text.splitlines() if line.strip()]
        if not lines:
            continue

        page_rows: list[list[str]] = []
        for line in lines:
            if "\t" in line:
                cells = [cell.strip() for cell in line.split("\t") if cell.strip()]
                has_table_structure = True
            else:
                cells = [cell.strip() for cell in re.split(r"\s{2,}", line) if cell.strip()]
                if len(cells) > 1:
                    has_table_structure = True

            if cells:
                page_rows.append(cells)

        if page_rows:
            extracted_rows.extend(page_rows)

    if not extracted_rows or not has_table_structure:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="PDF has no extractable table structure for Excel export.",
        )

    for row_index, row in enumerate(extracted_rows, start=1):
        for column_index, value in enumerate(row, start=1):
            default_sheet.cell(row=row_index, column=column_index, value=str(value))

    workbook.save(output_xlsx)
    workbook.close()


def convert_excel_to_pdf(input_xlsx: Path, output_pdf: Path) -> None:
    """Render a simple, readable XLSX workbook to a PDF."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="openpyxl is not installed") from exc

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="reportlab is not installed") from exc

    workbook = load_workbook(filename=input_xlsx, read_only=True, data_only=True)
    if not workbook.sheetnames:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel workbook is empty")

    pdf_canvas = canvas.Canvas(str(output_pdf), pagesize=letter)
    width, height = letter
    max_rows_per_page = 20
    max_columns = 8
    left_margin = 16 * mm
    top_margin = 18 * mm
    cell_width = (width - 2 * left_margin) / max_columns
    cell_height = 7 * mm

    try:
        for sheet_index, sheet_name in enumerate(workbook.sheetnames):
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            pdf_canvas.setTitle(f"{sheet_name} export")
            pdf_canvas.setFont("Helvetica-Bold", 12)
            pdf_canvas.drawString(left_margin, height - 18 * mm, f"Sheet: {sheet_name}")

            row_cursor = 0
            current_page_rows = 0
            for row_index, row in enumerate(rows):
                if row_index % max_rows_per_page == 0 and row_index > 0:
                    pdf_canvas.showPage()
                    pdf_canvas.setFont("Helvetica-Bold", 12)
                    pdf_canvas.drawString(left_margin, height - 18 * mm, f"Sheet: {sheet_name}")
                    current_page_rows = 0

                y = height - (28 * mm + current_page_rows * cell_height)
                for column_index, cell_value in enumerate(row[:max_columns]):
                    text = "" if cell_value is None else str(cell_value)
                    if len(text) > 20:
                        text = text[:20] + "..."
                    x = left_margin + column_index * cell_width
                    pdf_canvas.setFillColor(colors.black)
                    pdf_canvas.setStrokeColor(colors.grey)
                    pdf_canvas.rect(x, y - cell_height, cell_width, cell_height, fill=0, stroke=1)
                    pdf_canvas.setFont("Helvetica", 8)
                    pdf_canvas.drawString(x + 3, y - 3 * mm, text)
                current_page_rows += 1

            if sheet_index < len(workbook.sheetnames) - 1:
                pdf_canvas.showPage()
    finally:
        pdf_canvas.save()
        workbook.close()

def convert_pdf_to_images(pdf_bytes: bytes, output_directory: Path) -> list[Path]:
    """Render each page of a PDF into a PNG image."""
    try:
        import fitz
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PyMuPDF is not installed") from exc

    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        if document.page_count == 0:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded PDF has no pages")

        output_directory.mkdir(parents=True, exist_ok=True)
        matrix = fitz.Matrix(2, 2)
        output_paths: list[Path] = []

        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            output_path = output_directory / f"page_{page_index + 1}.png"
            pixmap.save(str(output_path))
            output_paths.append(output_path)

        return output_paths
    finally:
        document.close()


def merge_pdfs(input_paths: list[Path], output_path: Path) -> None:
    """Merge multiple PDF files into one output PDF."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="pypdf is not installed") from exc

    writer = PdfWriter()
    for input_path in input_paths:
        reader = PdfReader(str(input_path))
        if len(reader.pages) == 0:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="One of the PDF files is empty")
        for page in reader.pages:
            writer.add_page(page)

    with output_path.open("wb") as output_file:
        writer.write(output_file)


def split_pdf(input_path: Path, output_directory: Path) -> list[Path]:
    """Split a PDF into one file per page."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="pypdf is not installed") from exc

    reader = PdfReader(str(input_path))
    if len(reader.pages) == 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded PDF has no pages")

    output_directory.mkdir(parents=True, exist_ok=True)
    output_paths: list[Path] = []

    for page_index, page in enumerate(reader.pages, start=1):
        writer = PdfWriter()
        writer.add_page(page)
        output_path = output_directory / f"page_{page_index}.pdf"
        with output_path.open("wb") as output_file:
            writer.write(output_file)
        output_paths.append(output_path)

    return output_paths


def convert_image_to_pdf(image_bytes: bytes, output_path: Path) -> None:
    """Convert a single JPG/PNG image into a PDF."""
    try:
        from PIL import Image
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Pillow is not installed") from exc

    image = Image.open(BytesIO(image_bytes))
    if image.mode in ("RGBA", "LA", "P"):
        background = Image.new("RGB", image.size, (255, 255, 255))
        if image.mode == "P":
            image = image.convert("RGBA")
        background.paste(image, mask=image.split()[-1] if image.mode in ("RGBA", "LA") else None)
        image = background
    else:
        image = image.convert("RGB")

    image.save(str(output_path), "PDF", resolution=100.0)


def convert_images_to_pdf(image_bytes_list: list[bytes], output_path: Path) -> None:
    """Merge multiple images into a single PDF document."""
    try:
        from PIL import Image
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Pillow is not installed") from exc

    converted_images = []
    for image_bytes in image_bytes_list:
        image = Image.open(BytesIO(image_bytes))
        if image.mode in ("RGBA", "LA", "P"):
            background = Image.new("RGB", image.size, (255, 255, 255))
            if image.mode == "P":
                image = image.convert("RGBA")
            background.paste(image, mask=image.split()[-1] if image.mode in ("RGBA", "LA") else None)
            image = background
        else:
            image = image.convert("RGB")
        converted_images.append(image)

    first_image, *remaining_images = converted_images
    first_image.save(str(output_path), "PDF", save_all=True, append_images=remaining_images, resolution=100.0)


def compress_pdf(input_path: Path, output_path: Path) -> None:
    """Compress a PDF using PyMuPDF optimization flags."""
    try:
        import fitz
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PyMuPDF is not installed") from exc

    document = fitz.open(str(input_path))
    if document.page_count == 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded PDF has no pages")

    try:
        document.save(str(output_path), garbage=4, deflate=True, clean=True, incremental=False)
    finally:
        document.close()
